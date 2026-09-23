"""
PartnerDatasetAdapter: loads, checks and windows partner collected data.

Pipeline for one dataset folder:

    dataset.json + recordings.csv          partner_schema.load_schema()
            |
    load each file (csv / xlsx / mat)      PartnerDatasetAdapter.load_file()
            |
    split long format files by joint       -> streams (run, asset, joint)
            |
    quality report                         PartnerDatasetAdapter.validate()
            |
    sliding windows -> features            PartnerDatasetAdapter.extract()

Data is real machine data from the partner's equipment (domain REAL_MACHINE).
It is never mixed with simulator telemetry at the raw level.
No Qt, no PyBullet, no ML.
"""

import csv
import hashlib
import os

import numpy as np

import config
from datasets.partner_features import (
    channel_feature_names, thermal_feature_names, thermal_features, window_features,
)
from datasets.partner_schema import load_schema

DOMAIN = "REAL_MACHINE"


class Stream(object):
    """Continuous signals of one asset/joint from one recording."""

    def __init__(self, recording, joint, times, signals):
        self.recording = recording
        self.joint = joint
        self.times = times
        self.signals = signals

    @property
    def group(self):
        """Leakage unit: windows of one run never go to both train and test."""
        return self.recording.run_id

    @property
    def stream_id(self):
        return "%s|%s|%s" % (self.recording.run_id, self.recording.asset_id, self.joint)


class PartnerDatasetAdapter(object):

    name = "PARTNER"

    def __init__(self, root=None):
        self.root = root or config.PARTNER_DATA_DIR
        self.schema = load_schema(self.root)

    # -- loading ------------------------------------------------------------

    def load_file(self, recording):
        """Returns (column_names, 2 D float array, joint values or None)."""
        fmt = self.schema.format
        if fmt == "csv":
            return self._load_csv(recording.path)
        if fmt == "xlsx":
            return self._load_xlsx(recording.path)
        return self._load_mat(recording.path)

    def _load_csv(self, path):
        with open(path, newline="", encoding="utf-8-sig") as handle:
            reader = csv.reader(handle, delimiter=self.schema.csv_delimiter)
            header = [h.strip() for h in next(reader)]
            rows = [row for row in reader if row and any(cell.strip() for cell in row)]
        return self._table(header, rows)

    def _load_xlsx(self, path):
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[self.schema.sheet] if self.schema.sheet else workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        body = [["" if v is None else v for v in row] for row in rows[1:]
                if row and any(v is not None for v in row)]
        return self._table(header, body)

    def _table(self, header, rows):
        joint_index = header.index(self.schema.joint_column) if self.schema.joint_column in header else None
        joints = [str(r[joint_index]).strip() for r in rows] if joint_index is not None else None
        numeric = np.full((len(rows), len(header)), np.nan)
        for j, name in enumerate(header):
            if j == joint_index:
                continue
            for i, row in enumerate(rows):
                if j < len(row):
                    try:
                        numeric[i, j] = float(row[j])
                    except (TypeError, ValueError):
                        numeric[i, j] = np.nan
        return header, numeric, joints

    def _load_mat(self, path):
        from scipy.io import loadmat
        data = loadmat(path)
        variable = self.schema.mat_variable
        if variable:
            if variable not in data:
                raise ValueError("variable %s not found in %s" % (variable, path))
            matrix = np.asarray(data[variable], dtype=np.float64)
            if matrix.ndim == 1:
                matrix = matrix.reshape(-1, 1)
            columns = [c.column for c in self.schema.channels]
            if self.schema.time_column:
                columns = [self.schema.time_column] + columns
            if matrix.shape[1] != len(columns) and matrix.shape[0] == len(columns):
                matrix = matrix.T
            if matrix.shape[1] != len(columns):
                raise ValueError("%s: matrix has %d columns, schema expects %d (%s)"
                                 % (path, matrix.shape[1], len(columns), ", ".join(columns)))
            return columns, matrix, None
        # One variable per column, named like the schema columns.
        columns = [c.column for c in self.schema.channels]
        if self.schema.time_column:
            columns = [self.schema.time_column] + columns
        arrays = []
        for column in columns:
            if column not in data:
                raise ValueError("variable %s not found in %s" % (column, path))
            arrays.append(np.asarray(data[column], dtype=np.float64).reshape(-1))
        length = min(a.shape[0] for a in arrays)
        return columns, np.column_stack([a[:length] for a in arrays]), None

    def streams(self, recording):
        header, table, joints = self.load_file(recording)
        missing = [c.column for c in self.schema.channels if c.column not in header]
        if missing:
            raise ValueError("%s: missing columns %s" % (recording.file, ", ".join(missing)))
        groups = {}
        if joints is None:
            groups["all"] = np.arange(table.shape[0])
        else:
            for i, joint in enumerate(joints):
                groups.setdefault(joint, []).append(i)
        result = []
        for joint, indices in sorted(groups.items()):
            rows = table[np.asarray(indices)]
            times = self._times(header, rows, recording)
            signals = dict((c.name, rows[:, header.index(c.column)] * c.scale) for c in self.schema.channels)
            result.append(Stream(recording, joint, times, signals))
        return result

    def _times(self, header, rows, recording):
        offset = recording.start_time_s or 0.0
        if self.schema.time_column and self.schema.time_column in header:
            t = rows[:, header.index(self.schema.time_column)]
            return t - (t[0] if np.isfinite(t[0]) else 0.0) + offset
        return offset + np.arange(rows.shape[0]) / float(self.schema.sample_rate_hz)

    # -- quality ------------------------------------------------------------

    def validate(self, progress=None):
        """Readable report of problems and warnings. Never raises for bad files."""
        schema = self.schema
        report = {"dataset": schema.to_dict(), "errors": [], "warnings": [], "recordings": []}
        runs = {}
        for index, recording in enumerate(schema.recordings):
            if progress:
                progress("checking %s" % recording.file, index, len(schema.recordings))
            entry = recording.to_dict()
            if not os.path.isfile(recording.path):
                report["errors"].append("%s: file not found" % recording.file)
                report["recordings"].append(entry)
                continue
            try:
                streams = self.streams(recording)
            except (IOError, OSError, ValueError, KeyError, IndexError) as exc:
                report["errors"].append("%s: %s" % (recording.file, exc))
                report["recordings"].append(entry)
                continue
            entry["streams"] = []
            for stream in streams:
                info = self._stream_quality(stream, report)
                entry["streams"].append(info)
            runs.setdefault(recording.run_id, []).append(recording)
            report["recordings"].append(entry)

        labels = {}
        for recording in schema.recordings:
            labels.setdefault(recording.label or "(none)", set()).add(recording.run_id)
        report["labels"] = dict((k, len(v)) for k, v in labels.items())
        report["runs"] = len(runs)
        if schema.label_mode == "fault_labels":
            if schema.healthy_label not in labels:
                report["warnings"].append("no recording labelled %s: anomaly detection needs healthy data"
                                          % schema.healthy_label)
            for label, run_ids in labels.items():
                if len(run_ids) < 2:
                    report["warnings"].append(
                        "label %s appears in only %d run: it cannot be tested on an unseen run"
                        % (label, len(run_ids)))
        if schema.label_mode == "run_to_failure":
            if len(runs) < 3:
                report["errors"].append("run_to_failure needs at least 3 complete runs, found %d" % len(runs))
            no_failure = [r for r, recs in runs.items() if all(x.failure_time_s is None for x in recs)]
            if no_failure:
                report["warnings"].append("runs without failure_time_s use their last sample as failure: %s"
                                          % ", ".join(sorted(no_failure)[:5]))
        if len(runs) < 2:
            report["warnings"].append("only one run_id: results can only be checked on later time "
                                      "blocks of the same run, which is weak evidence")
        report["ok"] = not report["errors"]
        return report

    def _stream_quality(self, stream, report):
        n = stream.times.shape[0]
        info = {"joint": stream.joint, "samples": int(n),
                "duration_s": float(stream.times[-1] - stream.times[0]) if n > 1 else 0.0}
        name = "%s joint %s" % (stream.recording.file, stream.joint)
        if n > 2:
            dt = np.diff(stream.times)
            if np.any(dt <= 0):
                report["errors"].append("%s: time is not strictly increasing" % name)
            rate = 1.0 / float(np.median(dt)) if np.median(dt) > 0 else 0.0
            info["measured_rate_hz"] = rate
            declared = self.schema.sample_rate_hz
            if declared and abs(rate - declared) > config.PARTNER_RATE_TOLERANCE * declared:
                report["warnings"].append("%s: timestamps give %.1f Hz, dataset.json says %.1f Hz"
                                          % (name, rate, declared))
        thermal = self.thermal_channels()
        for channel in thermal:
            values = stream.signals[channel]
            after = stream.times >= stream.times[0] + self.schema.warmup_skip_s
            if after.sum() > 10:
                t = stream.times[after]
                x = values[after]
                tc = t - t.mean()
                denominator = float(np.sum(tc * tc))
                rate = 60.0 * float(np.sum(tc * (x - x.mean())) / denominator) if denominator > 0 else 0.0
                if rate > config.PARTNER_WARMUP_RATE_C_PER_MIN:
                    report["warnings"].append(
                        "%s joint %s: %s still rising %.1f C/min after warmup_skip_seconds=%.0f. "
                        "Increase warmup_skip_seconds or record longer, otherwise warm-up looks "
                        "like a fault." % (stream.recording.file, stream.joint, channel, rate,
                                           self.schema.warmup_skip_s))
        if thermal and not self._named("ambient_temperature"):
            message = "no ambient_temperature channel: temperature is judged without ambient " \
                      "compensation, so a hot room looks like a fault"
            if message not in report["warnings"]:
                report["warnings"].append(message)
        for channel, values in stream.signals.items():
            fraction = float(np.mean(~np.isfinite(values))) if n else 1.0
            if fraction > config.PARTNER_MAX_NAN_FRACTION:
                report["warnings"].append("%s: channel %s has %.0f%% missing values"
                                          % (name, channel, 100 * fraction))
            if n and np.nanstd(values) == 0.0:
                report["warnings"].append("%s: channel %s is constant (disconnected sensor?)" % (name, channel))
        if info["duration_s"] < self.schema.window_s:
            report["warnings"].append("%s: shorter than one %.1f s window" % (name, self.schema.window_s))
        return info

    # -- features -----------------------------------------------------------

    def sample_rate(self, stream):
        if self.schema.sample_rate_hz:
            return float(self.schema.sample_rate_hz)
        dt = np.diff(stream.times)
        return 1.0 / float(np.median(dt))

    def channel_list(self):
        """Channels windowed the normal way (temperature is handled separately)."""
        channels = [(c.name, c.type) for c in self.schema.channels
                    if c.type not in ("temperature", "ambient_temperature")]
        types = dict((c.type, c.name) for c in self.schema.channels)
        if "position" in types and "target_position" in types:
            channels.append(("tracking_error", "tracking_error"))
        return channels

    def thermal_channels(self):
        return [c.name for c in self.schema.channels if c.type == "temperature"]

    def _named(self, channel_type):
        for c in self.schema.channels:
            if c.type == channel_type:
                return c.name
        return None

    def feature_names(self, rate):
        names = []
        for name, channel_type in self.channel_list():
            names.extend(channel_feature_names(name, channel_type, rate))
        ambient = self._named("ambient_temperature")
        power = self._named("power")
        if ambient:
            names.append("%s_mean" % ambient)
        for name in self.thermal_channels():
            names.extend(thermal_feature_names(name, ambient is not None, power is not None))
        return names

    def extract(self, progress=None, cancel=None):
        """
        All windows of all streams. Returns dict with
            X (windows, features), names, groups (run_id), streams (stream_id),
            labels, conditions, assets, joints, times (window end, s),
            life_fraction / rul_fraction (run_to_failure only, else NaN)
        Windows containing missing values are skipped and counted.
        """
        schema = self.schema
        rows, groups, stream_ids, labels, conditions, assets, joints, ends = [], [], [], [], [], [], [], []
        names = None
        skipped = 0
        run_streams = {}
        for index, recording in enumerate(schema.recordings):
            if progress:
                progress("features for %s" % recording.file, index, len(schema.recordings))
            if cancel is not None and cancel():
                raise KeyboardInterrupt("cancelled")
            for stream in self.streams(recording):
                rate = self.sample_rate(stream)
                stream_names = self.feature_names(rate)
                if names is None:
                    names = stream_names
                elif stream_names != names:
                    raise ValueError("%s: feature layout differs (different sample rate?)" % recording.file)
                signals = dict(stream.signals)
                types = dict((c.type, c.name) for c in schema.channels)
                if "tracking_error" in [c[0] for c in self.channel_list()]:
                    signals["tracking_error"] = signals[types["target_position"]] - signals[types["position"]]
                window = max(config.PARTNER_MIN_WINDOW_SAMPLES, int(round(schema.window_s * rate)))
                hop = max(1, int(round(schema.hop_s * rate)))
                channels = self.channel_list()
                thermal = self.thermal_channels()
                ambient_name = self._named("ambient_temperature")
                power_name = self._named("power")
                thermal_window = max(window, int(round(schema.temperature_window_s * rate)))
                thermal_start = dict((name, float(np.nanmean(signals[name][:window])))
                                     for name in thermal)
                warmup_end = stream.times[0] + schema.warmup_skip_s
                for start in range(0, stream.times.shape[0] - window + 1, hop):
                    if stream.times[start + window - 1] < warmup_end:
                        continue        # warm-up: excluded from training and scoring
                    part = dict((k, v[start:start + window]) for k, v in signals.items())
                    if any(not np.all(np.isfinite(v)) for v in part.values()):
                        skipped += 1
                        continue
                    t = stream.times[start:start + window]
                    values = window_features(part, t, channels, rate)
                    row = [v for _, v in values]
                    if ambient_name or thermal:
                        low = max(0, start + window - thermal_window)
                        slice_t = stream.times[low:start + window]
                        if ambient_name:
                            row.append(float(np.nanmean(signals[ambient_name][low:start + window])))
                        for name in thermal:
                            row.extend(thermal_features(
                                signals[name][low:start + window], slice_t, thermal_start[name],
                                signals[ambient_name][low:start + window] if ambient_name else None,
                                signals[power_name][low:start + window] if power_name else None))
                    rows.append(row)
                    groups.append(recording.run_id)
                    stream_ids.append(stream.stream_id)
                    labels.append(self.stream_label(recording, stream))
                    # Healthy behaviour differs per joint (load, gear ratio), so
                    # the joint is part of the operating setup.
                    conditions.append("%s | joint %s" % (recording.condition, stream.joint))
                    assets.append(recording.asset_id)
                    joints.append(stream.joint)
                    ends.append(float(t[-1]))
                run_streams.setdefault(recording.run_id, []).append(recording)
        if not rows:
            raise ValueError("no complete windows: check window_seconds and missing values")

        ends = np.array(ends)
        groups_array = np.array(groups)
        life = np.full(ends.shape[0], np.nan)
        if schema.label_mode == "run_to_failure":
            for run_id, recordings in run_streams.items():
                mask = groups_array == run_id
                declared = [r.failure_time_s for r in recordings if r.failure_time_s is not None]
                failure = max(declared) if declared else float(ends[mask].max())
                start = float(ends[mask].min()) - schema.window_s
                life[mask] = np.clip((ends[mask] - start) / max(failure - start, 1e-9), 0.0, 1.0)
        return {
            "X": np.array(rows, dtype=np.float64), "names": names, "groups": groups_array,
            "streams": np.array(stream_ids), "labels": np.array(labels),
            "conditions": np.array(conditions), "assets": np.array(assets),
            "joints": np.array(joints), "times": ends,
            "life_fraction": life, "rul_fraction": 1.0 - life,
            "skipped_windows": skipped, "label_mode": schema.label_mode,
            "healthy_label": schema.healthy_label, "domain": DOMAIN,
            "schema_hash": self.schema_hash(),
        }

    def stream_label(self, recording, stream):
        if (self.schema.label_mode == "fault_labels" and recording.fault_joint
                and str(stream.joint) != recording.fault_joint):
            return self.schema.healthy_label
        return recording.label or ""

    def schema_hash(self):
        """Identifies the channel layout a model was trained on."""
        layout = [(c.name, c.type, c.unit) for c in self.schema.channels]
        text = repr((layout, self.schema.window_s, self.schema.sample_rate_hz))
        return hashlib.sha1(text.encode("utf-8")).hexdigest()[:12]
